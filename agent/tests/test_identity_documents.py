from __future__ import annotations

import base64
import io
import zipfile

import pytest
from docx import Document

from ollama_code import identity_documents as helper


def test_docx_generation_roundtrips_resume_text_in_memory():
    response = helper.handle({"protocol_version": 1, "action": "generate_docx", "title": "Professional resume", "sections": [
        {"heading": "Experience", "text": "Built customer support tools.\nImproved response times."},
        {"heading": "Education", "text": "Bachelor of Science"},
    ]})
    result = helper.handle({"protocol_version": 1, "action": "extract_docx", "data_base64": response["data_base64"]})
    assert result["ok"]
    assert "Professional resume" in result["text"]
    assert "Improved response times." in result["text"]
    assert "Bachelor of Science" in result["text"]


def test_docx_extraction_includes_tables_without_merged_duplicates():
    document = Document()
    document.add_paragraph("Candidate")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Merged work experience"
    buffer = io.BytesIO()
    document.save(buffer)
    text = helper.extract_docx(base64.b64encode(buffer.getvalue()).decode())
    assert text.count("Merged work experience") == 1


@pytest.mark.parametrize("payload", [None, {}, {"protocol_version": 2}, {"protocol_version": 1, "action": "shell"}])
def test_docx_protocol_is_closed(payload):
    with pytest.raises(helper.IdentityDocumentError):
        helper.handle(payload)


def test_docx_parser_errors_do_not_echo_source_contents():
    data = io.BytesIO()
    with zipfile.ZipFile(data, "w") as archive:
        archive.writestr("word/document.xml", "PRIVATE_SOURCE_BAD_XML")
    with pytest.raises(helper.IdentityDocumentError) as failure:
        helper.extract_docx(base64.b64encode(data.getvalue()).decode())
    assert "PRIVATE_SOURCE_BAD_XML" not in str(failure.value)


def test_docx_archive_rejects_entity_declarations():
    data = io.BytesIO()
    with zipfile.ZipFile(data, "w") as archive:
        archive.writestr("word/document.xml", '<!DOCTYPE root [<!ENTITY source SYSTEM "file:///private">]><root/>')
    with pytest.raises(helper.IdentityDocumentError, match="XML declarations"):
        helper.extract_docx(base64.b64encode(data.getvalue()).decode())


def test_docx_bounds_fail_before_document_output(monkeypatch):
    monkeypatch.setattr(helper, "MAX_TEXT_BYTES", 10)
    with pytest.raises(helper.IdentityDocumentError):
        helper.generate_docx("Resume", [{"heading": "Experience", "text": "too much document text"}])
    monkeypatch.setattr(helper, "MAX_SOURCE_BYTES", 5)
    with pytest.raises(helper.IdentityDocumentError):
        helper.extract_docx(base64.b64encode(b"large source").decode())


def test_docx_generation_rejects_xml_controls():
    with pytest.raises(helper.IdentityDocumentError, match="control"):
        helper.generate_docx("Resume", [{"heading": "", "text": "bad\x01text"}])
