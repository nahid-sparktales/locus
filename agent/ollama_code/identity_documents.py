"""Private, memory-only DOCX extraction and generation for Identity Vault.

The native vault owns file access, encryption, consent and process deadlines.
This helper never writes source text, output documents or parser diagnostics to
disk. Its sole stdout record is a bounded versioned JSON response.
"""
from __future__ import annotations

import base64
import binascii
import io
import json
import sys
import zipfile
from typing import Any

MAX_SOURCE_BYTES = 100 * 1024 * 1024
MAX_TEXT_BYTES = 5 * 1024 * 1024
MAX_WIRE_BYTES = 150 * 1024 * 1024
MAX_EXPANDED_BYTES = 100 * 1024 * 1024


class IdentityDocumentError(ValueError):
    pass


def _text(value: Any, limit: int = MAX_TEXT_BYTES) -> str:
    if not isinstance(value, str) or len(value.encode("utf-8")) > limit:
        raise IdentityDocumentError("Document text exceeds its supported limit.")
    # XML 1.0 cannot encode these controls; fail before constructing a document.
    if any(ord(char) < 32 and char not in "\t\r\n" for char in value):
        raise IdentityDocumentError("Document text contains unsupported control characters.")
    return value


def extract_docx(encoded: Any) -> str:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    if not isinstance(encoded, str) or len(encoded) > (MAX_SOURCE_BYTES + 2) // 3 * 4:
        raise IdentityDocumentError("Choose a DOCX document no larger than 100 MB.")
    try:
        data = base64.b64decode(encoded, validate=True)
    except (binascii.Error, ValueError) as exc:
        raise IdentityDocumentError("The document encoding is invalid.") from exc
    if len(data) > MAX_SOURCE_BYTES:
        raise IdentityDocumentError("Choose a DOCX document no larger than 100 MB.")
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > 10_000 or sum(item.file_size for item in entries) > MAX_EXPANDED_BYTES:
                raise IdentityDocumentError("The expanded document is too large.")
            if len({item.filename for item in entries}) != len(entries):
                raise IdentityDocumentError("The document contains duplicate archive entries.")
            if any(item.flag_bits & 1 for item in entries):
                raise IdentityDocumentError("Password-protected documents are not supported.")
            if "word/document.xml" not in archive.namelist():
                raise IdentityDocumentError("Choose a valid DOCX document.")
            for item in entries:
                if item.filename.lower().endswith((".xml", ".rels")):
                    xml = archive.read(item)
                    if b"<!DOCTYPE" in xml.upper() or b"<!ENTITY" in xml.upper():
                        raise IdentityDocumentError("The document contains unsupported XML declarations.")
        document = Document(io.BytesIO(data))
        parts: list[str] = []
        size = 0

        def collect(container: Any) -> None:
            nonlocal size
            for block in container.iter_inner_content():
                if isinstance(block, Paragraph):
                    text = block.text.strip()
                    if not text:
                        continue
                    size += len(text.encode("utf-8")) + 1
                    if size > MAX_TEXT_BYTES:
                        raise IdentityDocumentError("Extracted text exceeds the 5 MB limit.")
                    parts.append(text)
                elif isinstance(block, Table):
                    seen: set[Any] = set()
                    for row in block.rows:
                        for cell in row.cells:
                            if cell._tc not in seen:
                                seen.add(cell._tc)
                                collect(cell)

        collect(document)
        return "\n".join(parts)
    except IdentityDocumentError:
        raise
    except Exception as exc:
        # Parser exception strings can include source content or XML fragments.
        raise IdentityDocumentError("The DOCX document could not be read.") from exc


def generate_docx(title: Any, sections: Any) -> str:
    from docx import Document
    from docx.shared import Inches, Pt

    title = _text(title or "", 1_000)
    if not isinstance(sections, list) or not 1 <= len(sections) <= 100:
        raise IdentityDocumentError("A draft needs between 1 and 100 sections.")
    validated: list[tuple[str, str]] = []
    size = len(title.encode())
    for section in sections:
        if not isinstance(section, dict):
            raise IdentityDocumentError("The draft section is invalid.")
        heading = _text(section.get("heading", ""), 1_000)
        text = _text(section.get("text", ""))
        size += len(heading.encode()) + len(text.encode())
        if size > MAX_TEXT_BYTES:
            raise IdentityDocumentError("Draft text exceeds the 5 MB limit.")
        validated.append((heading, text))
    document = Document()
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    document.core_properties.title = title
    document.styles["Normal"].font.name = "Calibri"
    document.styles["Normal"].font.size = Pt(11)
    for page in document.sections:
        page.top_margin = page.bottom_margin = Inches(.7)
        page.left_margin = page.right_margin = Inches(.8)
    if title:
        document.add_heading(title, 0)
    for heading, text in validated:
        if heading:
            document.add_heading(heading, 1)
        for paragraph in text.split("\n"):
            document.add_paragraph(paragraph)
    output = io.BytesIO()
    document.save(output)
    if output.tell() > MAX_SOURCE_BYTES:
        raise IdentityDocumentError("The generated document exceeds the supported limit.")
    return base64.b64encode(output.getvalue()).decode("ascii")


def handle(request: Any) -> dict[str, Any]:
    if not isinstance(request, dict) or request.get("protocol_version") != 1:
        raise IdentityDocumentError("Unsupported Identity Vault document protocol.")
    if request.get("action") == "extract_docx":
        return {"protocol_version": 1, "ok": True, "text": extract_docx(request.get("data_base64"))}
    if request.get("action") == "generate_docx":
        return {"protocol_version": 1, "ok": True, "data_base64": generate_docx(request.get("title"), request.get("sections"))}
    raise IdentityDocumentError("Unsupported Identity Vault document action.")


def main() -> None:
    try:
        raw = sys.stdin.buffer.readline(MAX_WIRE_BYTES + 1)
        if len(raw) > MAX_WIRE_BYTES:
            raise IdentityDocumentError("The document request is too large.")
        response = handle(json.loads(raw))
    except IdentityDocumentError as exc:
        response = {"protocol_version": 1, "ok": False, "error": str(exc)}
    except Exception:
        response = {"protocol_version": 1, "ok": False, "error": "The document request could not be processed."}
    print(json.dumps(response, ensure_ascii=True), flush=True)


if __name__ == "__main__":
    main()
